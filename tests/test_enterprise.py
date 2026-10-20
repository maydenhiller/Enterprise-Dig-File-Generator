"""Run against the real example files: set EXAMPLES to the folder holding them."""
import io, os, random, sys, zipfile
import pytest
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
import app

EX = os.environ.get("EXAMPLES", "/root/.claude/uploads/12d12e69-6887-5c44-990e-a999b131df49")

def find(part):
    # newest match wins, so a re-uploaded template replaces an older one
    hits = sorted((n for n in os.listdir(EX) if part in n),
                  key=lambda n: os.path.getmtime(os.path.join(EX, n)), reverse=True)
    if hits:
        return open(os.path.join(EX, hits[0]), "rb").read()
    pytest.skip(f"{part} not available")

@pytest.fixture(scope="module")
def digs():
    a = app.parse_dig_package(find("AID_115"), "AID_115.xlsm")
    b = app.parse_dig_package(find("AID_116"), "AID_116.xlsm")
    return a, b

def test_counts(digs):
    assert len(digs[0]) == 59 and len(digs[1]) == 63

def test_dig_02b_fields(digs):
    d = next(x for x in digs[0] if x.number == "02B")
    assert (d.odometer, d.station_ft) == (57975.75, 332977)
    assert d.upstream_reference == "AGM 090" and d.downstream_reference == "AGM 100"
    assert round(d.us_weld_distance, 2) == 9.35 and round(d.ds_weld_distance, 2) == 48.02
    assert d.county_state == "Morris County, KS" and d.tract_number == "5-K-MO-30"
    assert d.nearest_reference_label == "AGM 100 1717.92'"

def test_kmz_matches_by_aid_and_feature(digs):
    k = app.parse_kmz(find("KMZ"), "x.kmz")
    assert len(k.lines) == 1 and len(k.digs) == 129
    alld = digs[0] + digs[1]
    app.reconcile_with_kmz(alld, k)
    assert not [w for d in alld for w in d.warnings]

def test_report_and_cheat_sheet(digs):
    import openpyxl
    tpl = find("Autopopulates")
    d = next(x for x in digs[1] if x.number == "05A")
    out, warn = app.build_staking_report(tpl, d, app.ReportSettings("Hayden Miller"),
                                         app.read_template_info(tpl).phones)
    assert not warn
    wb = openpyxl.load_workbook(io.BytesIO(out), data_only=True)
    sr = wb["StakingReport"]
    assert sr["G4"].value == "Dig # 05A" and sr["C12"].value == "580-585-7118"
    assert sr["J5"].value == d.station_ft and sr["C24"].value == d.hca
    z = zipfile.ZipFile(io.BytesIO(out))
    assert "xl/vbaProject.bin" in z.namelist()
    # pre-dig photo placeholders must keep their own media part
    assert "xl/media/image6.jpeg" in z.namelist()
    ch = app.build_cheat_sheet(find("for_ascending"), digs[0], auto_notes=True)
    ws = openpyxl.load_workbook(io.BytesIO(ch)).worksheets[0]
    assert ws["A6"].value.endswith("U/S Weld") and ws["B6"].value == "=B7-D6"
    first = min(digs[0], key=lambda x: x.odometer)
    assert (ws["I7"].value, ws["J7"].value) == (first.latitude, first.longitude)
    assert ws["I6"].value is None

def test_aerial_exact_span():
    z, f, mpp = app.aerial_geometry(38.7, 2600, 1600, 256, 19)
    assert abs(mpp * 1600 * 3.28084 - 2600) < 1
    assert 1.0 <= f < 2.0


def test_title_line_parsing():
    d = app.Dig(package_id="115", number="03A",
                header='Line ID 600-601-602 / Asmt ID 115, East Leg Mainline, 8.625" MP 52 to Kearney')
    i = app.asset_info(d)
    assert (i.line_id, i.assessment_id, i.line_name, i.segment_name, i.dig_number) == (
        "600-601-602", "115", "East Leg Mainline", '8.625" MP 52 to Kearney', "03A")
    d2 = app.Dig(package_id="128", number="59A",
                 header='Line ID 621 / Asmt ID 128,  Clinton Lateral Loop, 10.75" 10"  Iowa City to Clinton')
    j = app.asset_info(d2)
    assert (j.line_id, j.line_name, j.segment_name) == (
        "621", "Clinton Lateral Loop", '10.75" 10" Iowa City to Clinton')


def test_output_names(digs):
    d = next(x for x in digs[0] if x.number == "02B")
    assert app.staking_report_name(d) == "AID 115/Staking Reports/02-03 Survey Staking Report_Dig#02B.xlsm"
    assert app.excavation_report_name(d) == (
        "AID 115/Excavation Survey Reports/ExcavationSurvey Report_Dig 02B.pdf")
    assert app.profile_report_name(d) == "AID 115/Profile Reports/Profile Dig #02B.xlsx"


def test_excavation_report_is_an_incremental_update(digs):
    import re
    from pypdf import PdfReader
    tpl = find("ExcavationSurveyReport_Dig_XX")
    d = next(x for x in digs[1] if x.number == "02A")
    out, warn = app.build_excavation_report(tpl, d)
    assert not warn and out.startswith(tpl)          # original (signed) bytes untouched
    _, _, xml = app._xfa_datasets(PdfReader(io.BytesIO(out)))
    data = xml.split("<dd:dataDescription")[0]
    def val(name):
        m = re.search(r"<%s\s*>([^<]*)</%s\s*>" % (name, name), data)
        return m.group(1) if m else ""
    assert val("pipelineCode") == "763" and val("pipelineNumber") == "600-601-602"
    assert val("pipelineName") == "East Leg Mainline" and val("assessmentsegmentAssessmentID") == "116"
    assert val("assessmentsegmentName") == "8.625&quot; Kearney to Moberly" or \
        val("assessmentsegmentName") == '8.625" Kearney to Moberly'
    assert val("excavationNumber") == "02A" and val("excavationSurveyReportComment") == ""
    assert val("utilityLocationAccuracy") == "Sub-Meter"
    assert "<pipeconnectorNumber/>" in data or "<pipeconnectorNumber\n/>" in data


def test_profile_report(digs):
    import openpyxl
    tpl = find("Enterprise-Profile_Dig_XX")
    d = next(x for x in digs[0] if x.number == "03A")
    out, warn = app.build_profile_report(tpl, d)
    assert not warn
    wb = openpyxl.load_workbook(io.BytesIO(out))
    assert wb.sheetnames == ["Profile"]
    ws = wb.active
    assert ws["E3"].value == "East Leg Mainline" and ws["E4"].value == '8.625" MP 52 to Kearney'
    assert ws["E5"].value == "Line ID 600-601-602 /ASMT ID 115" and ws["C9"].value == "03A"
    assert ws["M7"].value is None and ws["E2"].value == "Enterprise Products"


def test_zip_layout_and_lazy_pdf(digs):
    import tempfile
    tpl = find("ExcavationSurveyReport_Dig_XX")
    d = digs[0][0]
    eager, _ = app.build_excavation_report(tpl, d)
    lazy, _ = app.build_excavation_report(tpl, d, lazy=True)
    # report IDs are random, so compare everything but that one GUID
    assert len(lazy) == len(eager) and lazy.template is tpl
    files = {app.excavation_report_name(d): lazy, app.profile_report_name(d): b"x"}
    with tempfile.TemporaryFile() as handle:
        app.write_zip(files, handle)
        handle.seek(0)
        archive = zipfile.ZipFile(handle)
        assert sorted(archive.namelist()) == sorted(files)
        pdf = archive.read(app.excavation_report_name(d))
        assert pdf.startswith(tpl) and len(pdf) == len(eager)


def test_line_colours_come_from_kmz():
    from PIL import Image
    from types import SimpleNamespace as NS
    k = app.parse_kmz(find("6710_77_79_6711"), "x.kmz")
    assert len(k.lines) == 2
    assert k.line_colours == [(255, 0, 0), (255, 255, 0)]      # loop red, mainline yellow

    real = app._basemap_image
    app._basemap_image = lambda view, url, token: (Image.new("RGB", (view.width, view.height), (60, 70, 50)), 1)
    try:
        dig = NS(latitude=38.76151, longitude=-97.51396, name="Dig # 05A",
                 upstream_reference=None, downstream_reference=None,
                 us_ref_lat=None, us_ref_lon=None, ds_ref_lat=None, ds_ref_lon=None,
                 package_id="6710", feature_id="77001033")
        jpeg = app.render_aerial(dig, k, width=800, slot_aspect=1600 / 835)
    finally:
        app._basemap_image = real
    img = Image.open(io.BytesIO(jpeg)).convert("RGB")
    pixels = list(img.getdata())
    assert any(r > 200 and g < 60 and b < 60 for r, g, b in pixels)      # red line drawn
    assert any(r > 200 and g > 200 and b < 90 for r, g, b in pixels)     # yellow line drawn


def test_long_text_is_fitted_to_its_boxes(digs):
    import openpyxl
    tpl = find("Autopopulates")
    d = next(x for x in digs[1] if x.number == "05A")
    short = "Turn left on Main Street."
    long = ("Drive north on the highway and continue traveling for several miles. " * 12).strip()
    assert app.fit_font_size(short, 500, 100) == 11.0
    assert 6.0 <= app.fit_font_size(long, 500, 100) < 11.0
    d.directions = long
    d.cells["C41"] = "A-VERY-LONG-ALIGNMENT-SHEET-NAME-THAT-WILL-NOT-FIT-IN-ONE-LINE.pdf"
    out, warn = app.build_staking_report(tpl, d, app.ReportSettings("Hayden Miller"),
                                         app.read_template_info(tpl).phones)
    assert not warn
    z = zipfile.ZipFile(io.BytesIO(out))
    styles = z.read("xl/styles.xml").decode()
    assert 'shrinkToFit="1"' in styles
    sheet = z.read("xl/worksheets/sheet1.xml").decode()
    import re
    def style_of(ref):
        return int(re.search(rf'<c r="{ref}"[^>]*? s="(\d+)"', sheet).group(1))
    xfs = re.findall(r"<xf\b[^>]*?/>|<xf\b[^>]*?>.*?</xf>",
                     re.search(r"<cellXfs\b.*?</cellXfs>", styles, re.S).group(0), re.S)
    assert "shrinkToFit" in xfs[style_of("C23")]                       # one-line cell
    fonts = re.findall(r"<font\b[^>]*?/>|<font\b[^>]*?>.*?</font>",
                       re.search(r"<fonts\b.*?</fonts>", styles, re.S).group(0), re.S)
    font_id = int(re.search(r'fontId="(\d+)"', xfs[style_of("A28")]).group(1))
    assert float(re.search(r'<sz val="([\d.]+)"', fonts[font_id]).group(1)) < 11   # wrapped box
    openpyxl.load_workbook(io.BytesIO(out))                             # still a valid workbook


def test_legend_does_not_cover_the_dig():
    from PIL import Image
    from types import SimpleNamespace as NS
    lat, lon = 29.936451, -94.297241
    pipeline = app.PipelineData()
    for i in range(7):
        offset = (i - 3) * 0.00022
        pipeline.lines.append([(lon - 0.01, lat + offset), (lon + 0.01, lat + offset)])
        pipeline.line_names.append(f"A fairly long pipeline name number {i}")
        pipeline.line_colours.append((255, 0, 0))
    dig = NS(latitude=lat, longitude=lon, name="Dig # 06B", upstream_reference=None,
             downstream_reference=None, us_ref_lat=None, us_ref_lon=None,
             ds_ref_lat=None, ds_ref_lon=None, package_id="1", feature_id="1")
    real = app._basemap_image
    app._basemap_image = lambda view, url, token: (Image.new("RGB", (view.width, view.height), (60, 80, 50)), 1)
    try:
        jpeg = app.render_aerial(dig, pipeline, width=1600, slot_aspect=1600 / 835)
    finally:
        app._basemap_image = real
    img = Image.open(io.BytesIO(jpeg)).convert("RGB")
    width, height = img.size
    # The dig pin sits at the centre and its yellow must still be visible there.
    r, g, b = img.getpixel((width // 2, height // 2 - 52))
    assert r > 200 and g > 170 and b < 80
    # The legend takes a modest part of the picture, not most of it.
    light = sum(1 for px in img.crop((0, 0, width, height)).resize((160, 84)).getdata()
                if min(px) > 220)
    assert light < 160 * 84 * 0.25
